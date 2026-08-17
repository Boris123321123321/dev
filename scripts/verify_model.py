#!/usr/bin/env python3
"""Проверка пересчитанной книги: значения ячеек против независимого расчёта.

Зелёный пересчёт доказывает только то, что формулы вычислимы. Этот скрипт
проверяет, что они вычисляют ТО, ЧТО НУЖНО: считает ключевые показатели на
чистом Python из тех же входных чисел и сравнивает с тем, что лежит в книге.

Запускать ПОСЛЕ пересчёта формул (иначе openpyxl вернёт None).

    python3 scripts/verify_model.py output/eco_tiers_unit_economics.xlsx
"""

import argparse
import sys

from openpyxl import load_workbook

TOL = 0.01  # относительный допуск, 1%


def close(a, b):
    if a is None or b is None:
        return False
    if abs(b) < 1e-9:
        return abs(a) < 1e-6
    return abs(a - b) / abs(b) < TOL


def main():
    p = argparse.ArgumentParser()
    p.add_argument("path")
    a = p.parse_args()

    wb = load_workbook(a.path, data_only=True)
    assum, rev = wb["Допущения"], wb["Выручка_вход"]
    unit, grid = wb["Unit-экономика"], wb["Тир_Период"]
    migr, scen = wb["Миграция"], wb["Сценарии"]

    # --- входные числа, как их видит книга
    gm = assum["B14"].value
    base = assum["B15"].value
    months_window = assum["B17"].value
    mults = [assum[f"D{6+i}"].value for i in range(4)]
    users = [rev[f"{c}6"].value for c in "CDEF"]
    revenue = [rev[f"{c}6"].value for c in "GHIJ"]

    print(f"вход: ставка маржи={gm}, база бонуса={base:.4f} ₽/юзер/мес, "
          f"окно выгрузки={months_window} мес")
    print(f"      множители={mults}")
    print(f"      юзеров={users}")
    print(f"      ВВ/юзер={revenue}\n")

    checks = []

    def chk(name, got, want):
        checks.append((name, got, want, close(got, want)))

    # --- окно выгрузки в месяцах (было ×365/90, стало формула по датам)
    chk("Допущения: месяцев в окне = 14", months_window, 14)

    # --- unit-экономика по тирам
    for i, cl in enumerate("CDEF"):
        tier = unit[f"{cl}4"].value
        chk(f"Unit {tier}: грязная маржа/юзер/мес",
            unit[f"{cl}8"].value, revenue[i] * gm)
        chk(f"Unit {tier}: бонус с механикой/юзер/мес",
            unit[f"{cl}10"].value, base * mults[i])
        chk(f"Unit {tier}: доплата/юзер/мес",
            unit[f"{cl}11"].value, base * (mults[i] - 1))
        chk(f"Unit {tier}: вклад/юзер/мес",
            unit[f"{cl}12"].value, revenue[i] * gm - base * mults[i])

    # средневзвешенный множитель не должен быть циклической ссылкой
    chk("Unit: средневзв. множитель по когорте",
        unit["G5"].value,
        sum(u * m for u, m in zip(users, mults)) / sum(users))

    # --- годовые итоги по когорте (лист Тир_Период, столбец N блока ИТОГО)
    # находим строки блока ИТОГО по подписям в колонке A
    labels = {}
    for r in range(1, grid.max_row + 1):
        v = grid[f"A{r}"].value
        if isinstance(v, str):
            labels.setdefault(v, []).append(r)

    def total_row(label):
        return labels[label][-1]  # последнее вхождение = блок ИТОГО

    year_rev = sum(u * rv for u, rv in zip(users, revenue)) * 12
    year_topup = sum(u * base * (m - 1) for u, m in zip(users, mults)) * 12
    year_base = sum(users) * base * 12

    chk("Тир_Период ИТОГО: ВВ за 2026",
        grid[f"N{total_row('ДОХОД: валовая выручка, ₽')}"].value, year_rev)
    chk("Тир_Период ИТОГО: грязная маржа за 2026",
        grid[f"N{total_row('ДОХОД: грязная маржа, ₽')}"].value, year_rev * gm)
    chk("Тир_Период ИТОГО: базовый бонус за 2026",
        grid[f"N{total_row('РАСХОД: эко-бонус базовый, ₽')}"].value, year_base)
    chk("Тир_Период ИТОГО: слепая доплата за 2026",
        grid[f"N{total_row('РАСХОД: доплата от механики, ₽')}"].value, year_topup)
    chk("Тир_Период ИТОГО: доплата % от маржи",
        grid[f"N{total_row('Доплата, % от грязной маржи')}"].value,
        year_topup / (year_rev * gm))

    # --- миграция: unit-экономика перехода
    for s, cl in enumerate("CDE"):
        chk(f"Миграция шаг {s+1}: Δ грязная маржа",
            migr[f"{cl}8"].value, (revenue[s + 1] - revenue[s]) * gm)
        chk(f"Миграция шаг {s+1}: Δ эко-бонус",
            migr[f"{cl}11"].value, base * (mults[s + 1] - mults[s]))
        chk(f"Миграция шаг {s+1}: чистый эффект на мигранта",
            migr[f"{cl}12"].value,
            (revenue[s + 1] - revenue[s]) * gm - base * (mults[s + 1] - mults[s]))
        chk(f"Миграция шаг {s+1}: пул-источник", migr[f"{cl}14"].value, users[s])

    # --- сценарии и порог безубыточности
    for cl in "CDE":
        conv = [scen[f"{cl}{5+s}"].value for s in range(3)]
        mo = scen[f"{cl}10"].value
        pools = users[:3]
        dmarg = [(revenue[s + 1] - revenue[s]) * gm for s in range(3)]
        dbon = [base * (mults[s + 1] - mults[s]) for s in range(3)]
        chk(f"Сценарий {cl}: мигрантов/год",
            scen[f"{cl}9"].value, sum(p * c for p, c in zip(pools, conv)))
        chk(f"Сценарий {cl}: прирост маржи/год",
            scen[f"{cl}11"].value,
            mo * sum(p * c * d for p, c, d in zip(pools, conv, dmarg)))
        chk(f"Сценарий {cl}: бонусы мигрантов/год",
            scen[f"{cl}12"].value,
            mo * sum(p * c * d for p, c, d in zip(pools, conv, dbon)))
        chk(f"Сценарий {cl}: слепая доплата/год", scen[f"{cl}13"].value, year_topup)

    mo = scen["C10"].value
    denom = mo * sum(users[s] * ((revenue[s + 1] - revenue[s]) * gm
                                - base * (mults[s + 1] - mults[s])) for s in range(3))
    chk("Порог безубыточности", scen["C19"].value, year_topup / denom)

    # --- вывод
    bad = [c for c in checks if not c[3]]
    for name, got, want, ok in checks:
        mark = "OK  " if ok else "ОШИБКА"
        g = f"{got:,.4f}" if isinstance(got, (int, float)) else repr(got)
        w = f"{want:,.4f}" if isinstance(want, (int, float)) else repr(want)
        if not ok:
            print(f"{mark} {name}: в книге {g}, ожидалось {w}")
    print(f"\nпроверок: {len(checks)}, расхождений: {len(bad)}")

    if bad:
        sys.exit(1)
    print("все ключевые ячейки совпали с независимым расчётом")

    # справочные величины
    print(f"\nсправочно за 2026 год:")
    print(f"  ВВ когорты            {year_rev:>18,.0f} ₽")
    print(f"  грязная маржа         {year_rev*gm:>18,.0f} ₽")
    print(f"  эко-бонусы базовые    {year_base:>18,.0f} ₽")
    print(f"  слепая доплата        {year_topup:>18,.0f} ₽")
    print(f"  порог безубыточности  {scen['C19'].value*100:>18.4f} %")


if __name__ == "__main__":
    main()
